from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "CLARAYO_Project_Documentation.docx"
ASSETS = ROOT / "docs" / "generated"
ASSETS.mkdir(parents=True, exist_ok=True)
BLUE = "278CC2"
DARK = "243746"
PALE = "EAF7FD"
GRAY = "D9E2E8"

def font(size, bold=False):
    paths = [Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf")]
    for path in paths:
        if path.exists(): return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()

def rounded_box(draw, xy, title, subtitle="", fill="#EAF7FD", outline="#278CC2"):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    x1,y1,x2,y2=xy
    draw.text(((x1+x2)//2, y1+24), title, font=font(30, True), fill="#243746", anchor="ma")
    if subtitle:
        lines=[]
        words=subtitle.split(); line=""
        for word in words:
            test=(line+" "+word).strip()
            if draw.textlength(test,font=font(20))>x2-x1-35: lines.append(line); line=word
            else: line=test
        if line: lines.append(line)
        for i,line in enumerate(lines[:3]): draw.text(((x1+x2)//2,y1+70+i*25),line,font=font(20),fill="#506979",anchor="ma")

def arrow(draw, a, b):
    draw.line([a,b], fill="#278CC2", width=5)
    import math
    ang=math.atan2(b[1]-a[1],b[0]-a[0]); s=16
    p1=(b[0]-s*math.cos(ang-.5),b[1]-s*math.sin(ang-.5)); p2=(b[0]-s*math.cos(ang+.5),b[1]-s*math.sin(ang+.5))
    draw.polygon([b,p1,p2],fill="#278CC2")

def architecture(path):
    im=Image.new("RGB",(1800,780),"white"); d=ImageDraw.Draw(im)
    boxes=[(70,260,340,520),(450,260,750,520),(860,260,1170,520),(1280,260,1730,520)]
    rounded_box(d,boxes[0],"Users","Customer  Seller  Administrator")
    rounded_box(d,boxes[1],"React Client","Routes  Context  Axios  Recharts")
    rounded_box(d,boxes[2],"REST API","Express  JWT  Middleware  Validation")
    rounded_box(d,boxes[3],"MongoDB","Users  Stores  Products  Categories  Orders  Reviews")
    for i in range(3):
        a=((boxes[i][2]+boxes[i+1][0])//2-35,390); b=((boxes[i][2]+boxes[i+1][0])//2+35,390)
        arrow(d,a,b); arrow(d,(b[0],420),(a[0],420))
    d.text((900,80),"CLARAYO System Architecture",font=font(42,True),fill="#243746",anchor="ma")
    d.text((900,650),"HTTPS and JSON requests flow through role-protected endpoints; MongoDB remains the source of truth.",font=font(24),fill="#506979",anchor="ma")
    im.save(path)

def erd(path):
    im=Image.new("RGB",(1800,1150),"white"); d=ImageDraw.Draw(im)
    boxes={
      "User":(70,80,500,330),"Store":(685,70,1115,300),"Category":(1300,80,1730,300),
      "Product":(685,440,1115,760),"Order":(70,830,500,1070),"Review":(1300,840,1730,1070)}
    content={
      "User":"name  email  passwordHash\nrole  phone  address\nwishlist  accountStatus",
      "Store":"seller FK  storeName  slug\ndescription  logo  banner\napprovalStatus",
      "Category":"name  slug\ndescription  image",
      "Product":"category FK  seller FK  store FK\nname  price  stock  SKU\nstatus  rating  sold",
      "Order":"customer FK  items[]\nshippingAddress  totals\npayment and fulfillment status",
      "Review":"customer FK  product FK  order FK\nrating  comment"}
    for n,xy in boxes.items(): rounded_box(d,xy,n,content[n],"#F7FBFE","#278CC2")
    # connectors and cardinality labels
    links=[("User","Store","1","0..1"),("Store","Product","1","many"),("Category","Product","1","many"),("User","Product","1","many"),("User","Order","1","many"),("Order","Review","1","many"),("Product","Review","1","many")]
    centers={k:((v[0]+v[2])//2,(v[1]+v[3])//2) for k,v in boxes.items()}
    for a,b,ca,cb in links:
        pa,pb=centers[a],centers[b]; d.line([pa,pb],fill="#8CA4B3",width=3)
        d.text((pa[0]+12,pa[1]+8),ca,font=font(18,True),fill="#243746"); d.text((pb[0]+12,pb[1]+8),cb,font=font(18,True),fill="#243746")
    # redraw boxes above lines
    for n,xy in boxes.items(): rounded_box(d,xy,n,content[n],"#F7FBFE","#278CC2")
    im.save(path)

architecture(ASSETS/"architecture.png")
erd(ASSETS/"database_design.png")

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor.from_string(DARK)
styles['Normal'].paragraph_format.space_after=Pt(7); styles['Normal'].paragraph_format.line_spacing=1.15
for name,size in [('Title',30),('Heading 1',19),('Heading 2',14)]:
    s=styles[name]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0); s.paragraph_format.space_before=Pt(14); s.paragraph_format.space_after=Pt(7)
# Remove theme-provided paragraph borders from the title style.
title_ppr=styles['Title']._element.get_or_add_pPr()
for existing in title_ppr.findall(qn('w:pBdr')): title_ppr.remove(existing)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def borders(table):
    tblPr=table._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),GRAY); b.append(e)
    tblPr.append(b)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,BLUE); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],"F4F8FA")
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return t
def bullets(items):
    for item in items: doc.add_paragraph(item,style='List Bullet')
def number(items):
    for index,item in enumerate(items,1):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.first_line_indent=Inches(-.22)
        p.add_run(f'{index}.  ').bold=True; p.add_run(item)
def figure(path, caption, width=6.6):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(width))
    c=doc.add_paragraph(caption); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic=True; c.runs[0].font.size=Pt(9)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.space_after=Pt(14)
r=p.add_run('●'); r.font.color.rgb=RGBColor.from_string('69BDE5'); r.font.size=Pt(30)
t=doc.add_paragraph('CLARAYO Online Shopping System',style='Title'); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
s=doc.add_paragraph('Project Documentation'); s.alignment=WD_ALIGN_PARAGRAPH.CENTER; s.runs[0].font.size=Pt(18); s.runs[0].font.bold=True
doc.add_paragraph('\n')
lead=doc.add_paragraph('A full-stack, multi-vendor marketplace for customers, independent sellers, and marketplace administrators. This document describes the implemented system, its architecture and data model, representative interfaces, and verification results.')
lead.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n')
d=doc.add_paragraph('Prepared 23 October 2026'); d.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

doc.add_heading('Contents',0)
for i,title in enumerate(['Project Title','Introduction','Problem Statement','Objectives','Target Users','System Features','Technologies Used','System Architecture','Database Design','Screenshots','Testing','Conclusion'],1):
    doc.add_paragraph(f'{i}. {title}')
doc.add_page_break()

doc.add_heading('1 Project Title',1)
doc.add_paragraph('CLARAYO Online Shopping System')

doc.add_heading('2 Introduction',1)
doc.add_paragraph('CLARAYO is a web-based online shopping system that brings customers, independent sellers, and administrators into one marketplace. Customers can discover products, maintain a cart and wishlist, place Cash on Delivery orders, monitor fulfillment, and submit verified reviews. Sellers apply for a store, publish and maintain products, control inventory, and process the part of each order assigned to them. Administrators supervise users, seller applications, stores, products, categories, orders, and marketplace reports.')
doc.add_paragraph('The project uses a React single-page application and an Express REST API backed by MongoDB. Authentication uses signed JSON Web Tokens, and protected routes enforce customer, seller, and administrator responsibilities. The result is a working foundation for a multi-vendor marketplace with clear separation between presentation, business logic, and persistent data.')

doc.add_heading('3 Problem Statement',1)
doc.add_paragraph('Small sellers often rely on disconnected tools for listing products, monitoring stock, receiving orders, and communicating order progress. Customers must search across separate channels and may lack consistent product information or order tracking. At the same time, marketplace administrators need a controlled process for approving sellers and managing inappropriate or inactive accounts.')
doc.add_paragraph('CLARAYO addresses this problem by centralizing catalog management, customer ordering, seller fulfillment, and administrative oversight. It gives each role an appropriate workspace while keeping products, stores, orders, and account states synchronized in one database.')

doc.add_heading('4 Objectives',1)
doc.add_heading('4.1 General Objective',2)
doc.add_paragraph('To design and implement a secure, responsive, role-based online shopping system that supports the complete marketplace workflow from product discovery and checkout to seller fulfillment and administrator supervision.')
doc.add_heading('4.2 Specific Objectives',2)
number(['Provide secure registration, login, profile management, and password updates.','Allow customers to search, filter, save, purchase, and review products.','Support Cash on Delivery checkout with server-side price validation and stock control.','Allow prospective sellers to submit applications and approved sellers to manage stores, products, inventory, and orders.','Allow administrators to approve or reject seller applications and manage customers, sellers, categories, products, and orders.','Present role-specific statistics and reports in responsive dashboards.','Maintain referential relationships among users, stores, products, categories, orders, and reviews.'])

doc.add_heading('5 Target Users',1)
table(['User','Primary needs','Authorized activities'],[
['Customer','Convenient product discovery and reliable ordering','Register, sign in, edit profile, browse and filter products, manage cart and wishlist, place COD orders, track or cancel eligible orders, and review delivered purchases.'],
['Seller','A manageable storefront and fulfillment workspace','Apply for a store, maintain store information, create or edit products, manage inventory, process assigned orders, and view sales reports.'],
['Administrator','Marketplace governance and operational visibility','Review seller applications, manage users and stores, maintain categories, inspect products and orders, deactivate accounts or stores, and view marketplace analytics.']],[1.0,2.2,3.5])

doc.add_heading('6 System Features',1)
features=[
('Authentication and authorization','Registration, login, JWT sessions, protected routes, role checks, password hashing, profile updates, and account suspension.'),
('Product catalog','Search, category filtering, sorting, product details, stock visibility, featured listings, and seller attribution.'),
('Shopping tools','Persistent cart state, quantity controls, wishlist management, and a responsive customer interface.'),
('Checkout and orders','Cash on Delivery, server-calculated totals, unique order numbers, idempotent checkout, conditional stock updates, cancellation, and stock restoration.'),
('Seller workspace','Seller application, store profile, product CRUD, product status, stock management, order fulfillment, and sales reporting.'),
('Administrator workspace','Dashboard statistics, seller approvals, customer and store detail menus, deactivation, category CRUD, product/order oversight, and reports.'),
('Reviews','Verified reviews linked to a customer, product, and completed order, with aggregate product rating fields.'),
('Responsive interface','Desktop and mobile layouts, role-specific navigation, motion effects, toast notifications, icons, and charts.')]
table(['Feature group','Implemented capability'],features,[1.6,5.2])

doc.add_heading('7 Technologies Used',1)
table(['Layer','Technology','Purpose'],[
['Frontend','React 19, React Router, Vite','Component-based single-page application, routing, and optimized builds.'],
['Client data','Axios, React Context','REST communication plus authentication and cart state.'],
['Interface','CSS, Framer Motion, GSAP, Lucide React, Sonner, Recharts','Responsive styling, animation, icons, notifications, and dashboard charts.'],
['Backend','Node.js, Express 5','REST endpoints and application business logic.'],
['Security','JWT, bcryptjs, Helmet, CORS, express-rate-limit, express-validator','Authentication, password hashing, HTTP hardening, access control, throttling, and validation.'],
['Database','MongoDB, Mongoose','Document storage, schemas, validation, indexes, references, and transactions.'],
['Development','npm, concurrently, Git','Dependency management, concurrent local services, and source control.']],[1.1,2.2,3.5])

doc.add_heading('8 System Architecture',1)
doc.add_paragraph('The browser renders the React application and sends JSON requests through Axios. Public endpoints provide catalog information; protected endpoints include a bearer token. Express routes pass requests through security, authentication, validation, and role middleware before models read or update MongoDB. Responses use a consistent structure containing success, data, message, and optional pagination fields.')
figure(ASSETS/'architecture.png','Figure 1. High-level CLARAYO architecture',6.8)
doc.add_heading('8.1 Request Flow',2)
number(['A user performs an action in the React interface.','Axios sends an HTTP request to an endpoint under /api.','Express applies security middleware, validation, authentication, and role authorization.','A route performs business logic through Mongoose models.','MongoDB returns or persists the required records.','The API sends JSON to React, which updates the interface and displays feedback.'])

doc.add_heading('9 Database Design',1)
doc.add_paragraph('CLARAYO uses MongoDB collections with Mongoose schemas. ObjectId references express relationships while embedded order items preserve the product name, image, price, quantity, seller, and store at purchase time. This snapshot design protects historical order accuracy even when a product later changes.')
figure(ASSETS/'database_design.png','Figure 2. MongoDB collection and reference design',6.7)
table(['Collection','Key fields and constraints'],[
['User','Unique email; hashed password excluded by default; role is customer, seller, or admin; account status is active or suspended; wishlist references products.'],
['Store','One unique store per seller; unique slug; approval status is pending, approved, rejected, or suspended.'],
['Category','Unique name and slug; optional description and image.'],
['Product','Unique slug and SKU; references category, seller, and store; nonnegative price and integer stock; active, draft, or archived status.'],
['Order','Unique order number and optional unique idempotency key; customer reference; embedded item snapshots and per-seller fulfillment records.'],
['Review','References customer, product, and order; rating from 1 to 5; compound uniqueness prevents duplicate reviews for the same purchase.']],[1.2,5.6])

doc.add_heading('10 Screenshots',1)
doc.add_paragraph('The following captures show representative public pages from the implemented React application. Catalog content depends on the connected API and database; the captured product page therefore shows its loading placeholders when no live API data is available to the preview build.')
figure(ROOT/'docs/screenshots/home.png','Figure 3. CLARAYO home page',6.7)
figure(ROOT/'docs/screenshots/login.png','Figure 4. Customer sign-in page',6.7)
figure(ROOT/'docs/screenshots/products.png','Figure 5. Product catalog with search, category, and sorting controls',6.7)

doc.add_heading('11 Testing',1)
doc.add_paragraph('Testing combined automated build and syntax checks with repository-level verification of critical workflows. Tests that require a configured MongoDB instance should be executed in an integration environment, particularly transaction-based checkout.')
tests=[
['T01','Frontend production build','Run npm run build in frontend','Vite transforms and bundles the React application','Passed'],
['T02','Admin route syntax','Run node --check routes/admin.js','No JavaScript syntax errors','Passed'],
['T03','Customer registration','Submit valid unique name, email, and 8+ character password','Customer created and JWT returned','Implementation verified'],
['T04','Invalid login or suspended account','Use invalid credentials or a suspended user','401 for invalid credentials; 403 for suspended account','Implementation verified'],
['T05','Product discovery','Search, filter by category, and sort products','Only active products from approved stores are returned','Implementation verified'],
['T06','Idempotent checkout','Repeat checkout with the same Idempotency-Key','Original order is returned without duplicate stock reduction','Integration test required'],
['T07','Insufficient stock','Order a quantity greater than available stock','409 response and no partial stock change','Integration test required'],
['T08','Order cancellation','Cancel a processing order','Order becomes cancelled and stock is restored','Integration test required'],
['T09','Seller isolation','Seller A attempts to update Seller B product','No product is exposed or modified','Implementation verified'],
['T10','Seller approval','Administrator approves a pending store','Store becomes approved and user role becomes seller','Implementation verified'],
['T11','Store deactivation','Administrator deactivates an approved store','Status becomes suspended and its products disappear publicly','Passed by code path'],
['T12','Category editing','Administrator edits category name or description','PUT updates category and refreshed row shows the saved values','Passed by code path']]
table(['ID','Test case','Procedure','Expected result','Result'],tests,[.55,1.25,2.1,2.25,.95])
doc.add_heading('11.1 Recommended Acceptance Test Sequence',2)
number(['Start MongoDB as a replica set and configure the backend environment.','Create the initial administrator and start both frontend and backend services.','Register a customer, apply as a seller, approve the application, and sign in again.','Create categories and products, then complete customer browsing, cart, checkout, cancellation, and review workflows.','Verify seller-specific product and order isolation using two seller accounts.','Suspend a customer and a store, then confirm authentication and public catalog behavior.'])

doc.add_heading('12 Conclusion',1)
doc.add_paragraph('CLARAYO accomplishes the core requirements of a multi-vendor online shopping system. It separates customer, seller, and administrator responsibilities; supports catalog, ordering, fulfillment, and governance workflows; and persists marketplace data through structured MongoDB schemas. The implementation also applies important safeguards such as password hashing, signed tokens, route authorization, input validation, rate limiting, server-side price calculation, idempotent checkout, and seller-scoped queries.')
doc.add_paragraph('The project demonstrates how a React client, REST API, Express backend, and MongoDB database can work together as a coherent production-style application. The main learning outcomes include role-based system design, transactional order handling, responsive interface development, document-oriented data modeling, API security, and full-stack integration. Future development can add online payment gateways, delivery-provider integrations, automated test suites, image storage, email notifications, and deeper analytics.')

# Page footer
for section in doc.sections:
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run('CLARAYO Project Documentation'); run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string('738493')

doc.core_properties.title='CLARAYO Online Shopping System Project Documentation'
doc.core_properties.subject='Full-stack online shopping system documentation'
doc.core_properties.author='CLARAYO Project Team'
doc.save(OUT)
print(OUT)
